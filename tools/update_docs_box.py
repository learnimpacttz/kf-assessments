# -*- coding: utf-8 -*-
# Adds the suggestion box (Sanduku la maoni) to the guide (1.8 -> 1.9) and manual (0.6 -> 0.7): field sections and the HQ section.
import docx, os
B = os.path.expanduser('~/LearnImpact_ED/KiuFunza Analysis/KF4_Endline_2026_Build/')
FIELD = [('b', 'Sanduku la maoni (suggestion box): on your first page, and under More. Write an idea, a problem, good news or a worry and press Send. The message is locked on your phone and only the Executive Director can read it. The system does not store your name, code or region unless you choose to tick "I want the ED to know it is me". Do not write other people\'s names if you want to stay anonymous. For something urgent that cannot wait, tell your RC or HQ directly.')]
HQ = [('h2', 'Suggestion box (HQ)'),
 ('p', 'The field team can send anonymous messages through Sanduku la maoni. Each message is locked on the sender\'s phone with the Executive Director\'s key. Only the Executive Director can read them, with a private key file kept on his own device. The server, the systems owner, Joram, the backups and the hosting company hold only locked text and cannot open it. The box is not counted in the usage pages or the activity log, so nobody can see who opened it. If the key file is lost, the messages cannot be recovered. The full description, including what it cannot promise, is in the Suggestion Box Protocol.')]
def insert_after_heading(doc, prefix, blocks):
    ps = doc.paragraphs
    idx = [i for i, p in enumerate(ps) if p.style.name == 'Heading 1' and p.text.startswith(prefix)]
    if not idx: print('missing', prefix); return
    nxt = next((p for p in ps[idx[0] + 1:] if p.style.name == 'Heading 1'), None)
    for k, t in blocks:
        par = doc.add_paragraph(t, style={'h2': 'Heading 2', 'p': 'Normal', 'b': 'List Bullet'}[k])
        if nxt is not None: nxt._p.addprevious(par._p)
def apply(doc):
    insert_after_heading(doc, '3. For test admins', FIELD); insert_after_heading(doc, '4. For regional and assistant coordinators', FIELD); insert_after_heading(doc, '5. For HQ', HQ)
def bump(doc, a, b):
    for p in doc.paragraphs:
        if a in p.text:
            for r in p.runs: r.text = r.text.replace(a, b)
g = B + 'KF4_Field_Command_Centre_User_Guide.docx'; d = docx.Document(g); apply(d); bump(d, 'Version 1.8', 'Version 1.9'); d.save(g)
m = B + 'KF4_Field_Operations_Manual_INTERIM_v0.6.docx'; d = docx.Document(m); apply(d); bump(d, 'Version 0.6', 'Version 0.7')
for t in d.tables:
    if t.rows and 'ersion' in t.rows[0].cells[0].text:
        c = t.add_row().cells
        for cell, v in zip(c, ['0.7', '21 Oct 2026', 'Suggestion box (Sanduku la maoni): anonymous, locked on the sender\'s phone, read only by the Executive Director.']): cell.text = v
d.save(m); os.rename(m, B + 'KF4_Field_Operations_Manual_INTERIM_v0.7.docx')
print('ok')
