# -*- coding: utf-8 -*-
# Adds the sampling attendance edit to the HQ section of the guide (1.6 -> 1.7) and manual (0.4 -> 0.5).
import docx, os
B = os.path.expanduser('~/LearnImpact_ED/KiuFunza Analysis/KF4_Endline_2026_Build/')
BL = [('h2', 'Correcting the attendance on a sampling form (HQ)'),
 ('p', 'Data cleaning, Suggest a deletion or an edit, Edit records: choose the Sampling project, find the school, tick the form, choose Pupils present on the sampling day, type the right number (1 to 20) and write the reason. Only the attendance answer changes. It is allowed only when 20 or fewer pupils were present, because the form draws no sample then and every present pupil is tested. When more than 20 were present the random draw of 20 depends on the number, so the edit is refused and the systems owner arranges a new draw.'),
 ('b', 'The request goes through the same chain as other edits: suggested, initiated by Joram, applied by Mundo. The original record is saved first, the change is made through KoBo, the record is read back and checked, and the dashboard rebuilds the Sampling data. The target for that class follows the new number.')]
def insert(doc):
    ps = doc.paragraphs
    idx = [i for i, p in enumerate(ps) if p.style.name == 'Heading 1' and p.text.startswith('5. For HQ')]
    if not idx: print('missing HQ section'); return
    nxt = next((p for p in ps[idx[0] + 1:] if p.style.name == 'Heading 1'), None)
    for k, t in BL:
        par = doc.add_paragraph(t, style={'h2': 'Heading 2', 'p': 'Normal', 'b': 'List Bullet'}[k])
        if nxt is not None: nxt._p.addprevious(par._p)
def bump(doc, a, b):
    for p in doc.paragraphs:
        if a in p.text:
            for r in p.runs: r.text = r.text.replace(a, b).replace('20 October 2026', '21 October 2026')
g = B + 'KF4_Field_Command_Centre_User_Guide.docx'; d = docx.Document(g); insert(d); bump(d, 'Version 1.6', 'Version 1.7'); d.save(g)
m = B + 'KF4_Field_Operations_Manual_INTERIM_v0.4.docx'; d = docx.Document(m); insert(d); bump(d, 'Version 0.4', 'Version 0.5')
for t in d.tables:
    if t.rows and 'ersion' in t.rows[0].cells[0].text:
        c = t.add_row().cells
        for cell, v in zip(c, ['0.5', '21 Oct 2026', 'Sampling attendance can be corrected from Data cleaning (20 or fewer present).']): cell.text = v
d.save(m); os.rename(m, B + 'KF4_Field_Operations_Manual_INTERIM_v0.5.docx')
print('ok')
