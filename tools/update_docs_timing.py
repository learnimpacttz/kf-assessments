# -*- coding: utf-8 -*-
# Timing in queries (27 Oct 2026): guide 2.3 -> 2.4, manual 1.1 -> 1.2. A section for coordinators and one for HQ.
import docx, os
B = os.path.expanduser('~/LearnImpact_ED/KiuFunza Analysis/KF4_Endline_2026_Build/')
HQ = [('h2', 'Waiting at HQ, and overdue from the field, are shown separately'),
 ('p', 'On the Bridge, No reply from the field counts queries where the person has not answered after two working days. Waited more than two working days at HQ counts queries the field has already answered that are waiting for the first HQ review or final approval, and shows the age of the oldest. These are not overdue from the teams, so the two are never added together. The Bridge figures for queries closed per day and the median time to approval now count approved queries; before, they showed zero.')]
ALL = [('h2', 'The times on a query, and records sent after the test day'),
 ('p', 'A query can show three different times. Started and finished are the times on the phone when the test was done. Sent is the time the record reached KoBo (East Africa Time). After the test is the wait between finishing and sending. Every query that has times carries this explanation, and the details show the start and finish of each test and when it was sent.'),
 ('p', 'Tests are sent on the day they are done. A record sent on a later day is flagged from the first day after the test (before, only two or more days were flagged) as Sent after the test day. The query says in plain words when it was tested, when the testing finished on the phone, when the last record was sent, how long it waited, and, when it applies, how long before the first test at the next school that day it was sent. Records sent just before the next school begins are the pattern to explain.'),
 ('b', 'The person is asked one thing: why the records were not sent on the day of the test, and where the phone was when they were sent. Poor signal is a real problem, so the reply should say what was tried (a place with better signal, a hotspot, sending from home that evening) and what will be done so it does not repeat. Records sent after midnight count as the next day. It is a reason to explain, not a reason to leave data on the phone.'),
 ('b', 'The two working days to reply to a Sent after the test day query count from the day the records were sent, and never from before 27 Oct 2026, when the check began. A query that could not have existed earlier is not overdue the day it appears.'),
 ('b', 'The checks about phones (a phone at two schools very close in time, tests that ran at the same time, one name on two phones) use the start and finish times on the phone, not the time of sending. Their wording now says so.')]
def insert(doc, heading, items=None):
    ps = doc.paragraphs
    idx = [i for i, p in enumerate(ps) if p.style.name == 'Heading 1' and p.text.startswith(heading)]
    if not idx: print('missing', heading); return
    nxt = next((p for p in ps[idx[0] + 1:] if p.style.name == 'Heading 1'), None)
    for k, t in (items or ALL):
        par = doc.add_paragraph(t, style={'h2': 'Heading 2', 'p': 'Normal', 'b': 'List Bullet'}[k])
        if nxt is not None: nxt._p.addprevious(par._p)
def bump(doc, a, b):
    for p in doc.paragraphs:
        if a in p.text:
            for r in p.runs: r.text = r.text.replace(a, b)
g = B + 'KF4_Field_Command_Centre_User_Guide.docx'; d = docx.Document(g); insert(d, '4. For regional and assistant coordinators'); insert(d, '5. For HQ', HQ); bump(d, 'Version 2.3', 'Version 2.4'); d.save(g)
m = B + 'KF4_Field_Operations_Manual_INTERIM_v1.1.docx'; d = docx.Document(m); insert(d, '4. For regional and assistant coordinators'); insert(d, '5. For HQ', HQ); bump(d, 'Version 1.1', 'Version 1.2')
for t in d.tables:
    if t.rows and 'ersion' in t.rows[0].cells[0].text:
        c = t.add_row().cells
        for cell, v in zip(c, ['1.2', '27 Oct 2026', 'Queries: clear timing wording, a check for records sent the day after the test, a timing list on each query, and separate HQ-waiting and field-overdue counts on the Bridge.']): cell.text = v
d.save(m); os.rename(m, B + 'KF4_Field_Operations_Manual_INTERIM_v1.2.docx')
print('ok')
