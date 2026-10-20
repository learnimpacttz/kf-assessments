# -*- coding: utf-8 -*-
# Adds Team patterns and the Visit map to the guide (1.5 -> 1.6) and manual (v0.3 -> v0.4). Field sections get the map; the HQ section gets both.
import docx, os
B = os.path.expanduser('~/LearnImpact_ED/KiuFunza Analysis/KF4_Endline_2026_Build/')
MAP_FIELD = [('b', 'Visit map: every school in your region as a dot on a simple map. Green is complete, gold is in progress, grey is not started. Tap a dot to see the school, how many pupils are tested and the visit dates. The dots are school locations, not people. Find it under Plan (coordinators) or Team (test admins).')]
HQ = [
 ('h2', 'Team patterns (Checks)'),
 ('p', 'Team patterns shows which test administrators may need support, so Mundo and Joram can help early. For each person it compares pace, timing, location and data-quality checks with others in the same role, and only flags repeated gaps after at least 3 test days and 20 tests. A flagged person has plain reasons (for example many tests faster than normal, a first test later than peers, few tests a day, phone location missing or away from the school), repeated-day patterns (for example fast tests mostly on Tuesdays), and possible causes to check first.'),
 ('b', 'Top of the page: counts (need attention, worth a conversation, within the usual range, too early to tell) and the share of records that carry a phone location. Then the people to talk to, with day-by-day strips, followed by where everyone sits on each measure, which weekdays stand out, a region table and a table of everyone.'),
 ('b', 'Use it as a prompt for a supportive conversation, never a verdict. Time of day, travel, school size and phone problems explain many gaps. Concerns about the data itself still go through the Queries chain. It changes no record, no payment and sends no message.'),
 ('b', 'Add what you learn in the HQ note on each person. Everyone in HQ can open the page except Advisors, and each opening is recorded in the activity log.'),
 ('b', 'The Bridge adds a line to Needs you now when someone needs attention.'),
 ('h2', 'Visit map and location checks (HQ)'),
 ('p', 'The Visit map shows all 325 schools coloured by status. Filter by region, status, visited today or location check needed. Tap a school for its tests and, for HQ only, a location check: how many records carry a phone location and how many were recorded within 500 m of the school, and who tested there.'),
 ('b', 'Only counts are kept from the phone location (with a point, within 500 m, more than 2 km away). The points themselves are never stored or shown. Phone location drifts 100 to 300 m inside buildings, so judge patterns and not one reading.'),
 ('b', 'About a third of records carry a phone location (KoBo only saves it when the phone has a fix). A red ring means many tests were recorded away from the school: it can be a wrong school coordinate or a phone that lost location, so check before assuming anything.'),
 ('b', 'School points come from the school list (301 of 325). Two schools use a point from OpenStreetMap, shown as approximate. A comparison of our points with OpenStreetMap is in School_coordinate_check_2026-10-20.csv; the remaining schools without a point are in Songwe and need the TAMISEMI list or a field visit.'),
]
def insert_after_heading(doc, head_prefix, blocks):
    ps = doc.paragraphs
    idx = [i for i, p in enumerate(ps) if p.style.name == 'Heading 1' and p.text.startswith(head_prefix)]
    if not idx: print('missing', head_prefix); return
    nxt = next((p for p in ps[idx[0] + 1:] if p.style.name == 'Heading 1'), None)
    for k, t in blocks:
        par = doc.add_paragraph(t, style={'h2': 'Heading 2', 'p': 'Normal', 'b': 'List Bullet'}[k])
        if nxt is not None: nxt._p.addprevious(par._p)
def apply(doc):
    insert_after_heading(doc, '3. For test admins', MAP_FIELD)
    insert_after_heading(doc, '4. For regional and assistant coordinators', MAP_FIELD)
    insert_after_heading(doc, '5. For HQ', HQ)
def bump(doc, a, b):
    for p in doc.paragraphs:
        if a in p.text:
            for r in p.runs: r.text = r.text.replace(a, b).replace('18 October 2026', '20 October 2026')
g = B + 'KF4_Field_Command_Centre_User_Guide.docx'; d = docx.Document(g); apply(d); bump(d, 'Version 1.5', 'Version 1.6'); d.save(g)
m = B + 'KF4_Field_Operations_Manual_INTERIM_v0.3.docx'; d = docx.Document(m); apply(d); bump(d, 'Version 0.3', 'Version 0.4')
for t in d.tables:
    if t.rows and 'ersion' in t.rows[0].cells[0].text:
        c = t.add_row().cells
        for cell, v in zip(c, ['0.4', '20 Oct 2026', 'Team patterns (HQ) and the Visit map; phone location kept only as counts; Editors can update the Bridge operations list, milestones and regional plans.']): cell.text = v
d.save(m); os.rename(m, B + 'KF4_Field_Operations_Manual_INTERIM_v0.4.docx')
print('ok')
