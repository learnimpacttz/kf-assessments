# -*- coding: utf-8 -*-
# Adds the Data cleaning chain (who suggests, confirms, decides), what the screen shows, and what to do after an edit or deletion is applied.
# Guide 1.7 -> 1.8, manual 0.5 -> 0.6. HQ section only.
import docx, os
B = os.path.expanduser('~/LearnImpact_ED/KiuFunza Analysis/KF4_Endline_2026_Build/')
BL = [('h2', 'Data cleaning: who does what, and what to do after a change (HQ)'),
 ('p', 'Mundo, Joram and the regional coordinators can suggest a deletion or an edit. What Mundo or the field suggests goes to Joram, who confirms it. What Joram suggests goes straight to Mundo. What the Executive Director starts also goes straight to Mundo, and the Executive Director can also give the final decision on it. Mundo gives the final decision and carries it out. Whoever confirmed a request cannot also decide it. The same chain covers edits, for example correcting the attendance on a sampling form.'),
 ('b', 'A button on a query that says "(opens the form)" does not send anything. Tick the record, write the reason and press Suggest at the bottom: the form says "Nothing is sent until you press this button", shows any error under the button, and after sending shows a green "Sent" banner and takes you to the new request. A strip at the top counts what is waiting for Joram, waiting for Mundo, sent back and finished, and each request says what happens next. The person who suggested gets a phone alert when it is passed on and when it is done.'),
 ('b', 'After Mundo applies a change: (1) read the result on the card: Done, with the KoBo count for a deletion or "read back and matched" for an edit; if it says Stopped, do not retry, read the reason and tell the builder. (2) Wait about 10 minutes while the dashboard rebuilds the totals. (3) Open the same request card: "Checks cleared by this change" lists the checks the change cleared and who had replied; the replies are kept. (4) Open the school in Regions and check the number that was changed; for an attendance edit the target equals the new attendance. (5) The cleared check disappears from the Queries page by itself; nothing is closed by hand. Other checks for the school need their own reply. (6) Do not edit the record again in KoBo and do not start rebuilds. The original is saved under Download the saved copy on the card.'),
 ('b', 'The Change log: deletions and edits lists every approved change, newest first, with who suggested, who confirmed and who decided. Refused attempts are written to the activity log in Admin.')]
def insert(doc):
    ps = doc.paragraphs
    idx = [i for i, p in enumerate(ps) if p.style.name == 'Heading 1' and p.text.startswith('5. For HQ')]
    if not idx: print('missing HQ section'); return
    nxt = next((p for p in ps[idx[0] + 1:] if p.style.name == 'Heading 1'), None)
    for k, t in BL:
        par = doc.add_paragraph(t, style={'h2': 'Heading 2', 'p': 'Normal', 'b': 'List Bullet'}[k])
        if nxt is not None: nxt._p.addprevious(par._p)
def bump(doc, a, b):
    for p in doc.paragraphs:
        if a in p.text:
            for r in p.runs: r.text = r.text.replace(a, b)
g = B + 'KF4_Field_Command_Centre_User_Guide.docx'; d = docx.Document(g); insert(d); bump(d, 'Version 1.7', 'Version 1.8'); d.save(g)
m = B + 'KF4_Field_Operations_Manual_INTERIM_v0.5.docx'; d = docx.Document(m); insert(d); bump(d, 'Version 0.5', 'Version 0.6')
for t in d.tables:
    if t.rows and 'ersion' in t.rows[0].cells[0].text:
        c = t.add_row().cells
        for cell, v in zip(c, ['0.6', '21 Oct 2026', 'Data cleaning chain (Mundo and Joram suggest, Joram confirms, Mundo decides), clearer screens, checks cleared by a change, and what to do after a change is applied.']): cell.text = v
d.save(m); os.rename(m, B + 'KF4_Field_Operations_Manual_INTERIM_v0.6.docx')
print('ok')
