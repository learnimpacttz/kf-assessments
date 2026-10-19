# -*- coding: utf-8 -*-
# Adds the HQ Bridge section to the guide (1.4 -> 1.5) and manual (v0.2 -> v0.3). HQ editions only: field editions drop the HQ section.
import docx, os
B = os.path.expanduser('~/LearnImpact_ED/KiuFunza Analysis/KF4_Endline_2026_Build/')
S5 = [('h2', 'The Bridge (the HQ landing page)'),
('p', 'Today opens on the Bridge. From top to bottom: Field status (day of 35, status on track, watch or behind, one plain sentence); Headline numbers (pupils tested, schools complete, data clean rate against 95%, checks open, sent the same day, people active today); Needs you now; Where we are heading (cumulative pupils tested, the plan, the finish range, and the milestones including those after field work); Regions; Data cleanliness; Early warnings; Targets; Management update.'),
('b', 'Tap a headline number to open its detail in place, with buttons that go to the full page. Click a region row to open it, or a column heading to sort. The Period, Region and Partner filters apply to the numbers and the board; Save this view remembers a combination on this device.'),
('b', 'Practice schools are never counted. Finish dates follow this year\'s plan first: planned visits per day, plus whatever slipped, scaled by how well each region keeps its plan (a 90% prior that gives way to real visits). Last year\'s history is only a fallback for a region with no locked plan.'),
('b', 'Clean means no open check and no pending correction, as a share of all records. A record edited after sending counts as not clean until the edit is approved. Open checks are counted once per person, day and grade.'),
('b', 'Needs you now is ranked by urgency and personal to you: queries waiting for your review level, queries with no reply after 2 working days, field questions, requests for Joram or Mundo, late notices, KoBo quota and cleaning pace.'),
('b', 'Dates: field work closes 4 Dec, cleaning 25 Nov, payment calculations 5 Dec (Mundo, Michael supporting), reports and the report to management, board and funders 15 Dec. An HQ Admin can change dates and status in the Targets card. Field work closes after the cleaning deadline, so the Bridge warns that records sent after 25 Nov need a final sweep.'),
('b', 'Management update: the ED update (for Michael) is written from the numbers every Monday at 07:00 and emailed to Michael as a draft; edit it in the Bridge, Copy or Download it, and Save as the update for this week. Board, funder and government versions are aggregate only and are for after the program ends.'),
('b', 'Desks: each HQ person has a desk that sets how the Bridge and the menu are arranged. Coordination (field coordination and first reviews): needs, regions and checks first. Management (final approval, procurement, contracting and operations): needs, targets and the operations list first. Leadership: headline numbers, timeline, targets and the update, with details folded away. Systems (the builder): system health first, then everything. A desk changes the view only, never what a person may do; set it in Admin > People & access, and anyone can try another desk from the Desk menu on the Bridge.'),
('b', 'Operations list (Management desk): contracts, partner disbursements and other processing from HQ, each with an owner, due date and status. Items due within 7 days, or past due, appear in Needs you now. The partner disbursements follow section 7 of each partner contract: Disbursement 4 on attendance at the start of central endline training, and Disbursement 5 on the endline report (due 15 Dec, paid about 18 Dec, 5% tax withheld). Bank details are never stored here.'),
('b', 'Leadership desk: view only. Period includes a Dates option (from and to) that totals pupils tested in that range; the region board shows each region\'s predicted finish; Reports and analytics buttons open the reports, compare, data explorer, usage and narratives pages.'),
('b', 'Team messages (Coordination and Management desks): a card near the top with Questions, Direct messages and News buttons, and two dropdowns to message a regional team or a coordinator in one step.'),
('b', 'Pages opened from the Bridge show a Back to Bridge button at the top; panels have a close (x). On a phone, wide tables become stacked cards, and the filters sit behind a Filters button. Systems desk only: the View as menu also lists each HQ team member, so you can see exactly what they see (read-only).'),
('b', 'Usage Explorer (Usage & analytics > Explorer): pick a period (today to 90 days, or your own dates), then filter by who, region, partner, page and device. You get page views, people active, visits and time with a comparison to the previous period; activity by day; a day-by-hour heat map; pages used; a funnel from roster to regular use; region and partner tables; and a people table. Tap a page, partner or region to filter to it. Tick people, or press Select the quiet ones, to send them a message with a phone alert, and download the table as CSV. Page and hour detail starts building from 18 Oct 2026; earlier days show only totals from saved actions.'),
('b', 'Where the queries are: a card showing how many queries wait at each step (person replies, team review, first HQ review, final approval, approved). It is the first card for anyone in the review chain, and the step that is yours is outlined and says "you"; tap a step to open those queries. Reminders: HQ people with pending reviews, requests or field questions get a banner on every page and a phone alert at 08:30 and 15:00, each opening the work itself (tap the banner x to pause it for an hour).'),
('b', 'Query details: opening Details to find the records on a query now loads the full record detail from KoBo: pupil ID and number, test admin, test date, start and end time, minutes, when it was sent to KoBo and how long after the test, the phone, and notes such as test set, distance from school or changed after sending. The review steps sit in one row of four equal boxes.'),
('b', 'Advisor access: in Admin > People & access each HQ person has an Access setting (Full or Advisor) and an optional last day. Advisor is read only and hides messaging (direct messages and Write and schedule), field questions, HQ notes, People, Phones, Usage, Payments, Data cleaning, Pilot day, Practice and the Operations list. Data behind those pages is blocked on the server, not only hidden. Use it for consultants and other outside readers.'),
('b', 'The search box also takes commands: type queries, overdue, questions, targets, usage and so on to jump there. Account menu > Arrange this page moves or hides any Bridge card.')]
def insert(doc, blocks):
    ps = doc.paragraphs
    idx = [i for i, p in enumerate(ps) if p.style.name == 'Heading 1' and p.text.startswith('5. For HQ')]
    if not idx: print('missing HQ section'); return
    nxt = next((p for p in ps[idx[0] + 1:] if p.style.name == 'Heading 1'), None)
    for k, t in blocks:
        par = doc.add_paragraph(t, style={'h2': 'Heading 2', 'p': 'Normal', 'b': 'List Bullet'}[k])
        if nxt is not None: nxt._p.addprevious(par._p)
def bump(doc, a, b):
    for p in doc.paragraphs:
        if a in p.text:
            for r in p.runs: r.text = r.text.replace(a, b).replace('17 October 2026', '18 October 2026')
g = B + 'KF4_Field_Command_Centre_User_Guide.docx'; d = docx.Document(g); insert(d, S5); bump(d, 'Version 1.4', 'Version 1.5'); d.save(g)
m = B + 'KF4_Field_Operations_Manual_INTERIM_v0.2.docx'; d = docx.Document(m); insert(d, S5); bump(d, 'Version 0.2', 'Version 0.3')
for t in d.tables:
    if t.rows and 'ersion' in t.rows[0].cells[0].text:
        c = t.add_row().cells
        for cell, v in zip(c, ['0.3', '18 Oct 2026', 'HQ Bridge: new landing page with targets, timeline, region board, cleanliness funnel, early warnings and the Monday management update.']): cell.text = v
d.save(m)
os.rename(m, B + 'KF4_Field_Operations_Manual_INTERIM_v0.3.docx')
print('ok')
