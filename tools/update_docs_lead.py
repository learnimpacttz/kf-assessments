# -*- coding: utf-8 -*-
# Adds the Lead Coordinator (LC) role to the guide (2.1 -> 2.2) and manual (0.9 -> 1.0): the coordinators' section and the HQ section.
import docx, os
B = os.path.expanduser('~/LearnImpact_ED/KiuFunza Analysis/KF4_Endline_2026_Build/')
RC = [('h2', 'Lead Coordinator (LC)'),
 ('p', 'Each partner organization has one Lead Coordinator: the RC of one of its regions, who also follows every region of the organization. GEP: Tanga, Manyara, Lindi, Mtwara. KACODA: Shinyanga, Mara, Singida. FAWOCHIWE: Kigoma, Rukwa, Songwe. A Lead Coordinator stays the RC of their own region and works there exactly as before.'),
 ('b', 'A bar at the top of Region, Plan, My team, Visit map, Queries, Phones, Data explorer and School reports switches between the regions of the organization. The own region comes first and is marked "(yours)".'),
 ('b', 'In the other regions a Lead Coordinator can read everything an RC reads: schools and the school drawer, queries with their responses and reviews, questions, plans and notices, the team, phones and school reports. They cannot reply, review, send notices, change a plan, send messages or edit records there. Only that region\'s own coordinators do.'),
 ('b', 'The title "RC · Lead Coordinator (LC)" shows next to their name on the All teams page, in the people lists and in their account menu.')]
HQ = [('b', 'Lead Coordinators (Admin, People and access): for each organization, choose which RC is the Lead Coordinator. It must be an active RC of a region of that organization. HQ can change or clear it at any time; the person sees the change the next time they open the app. Access is read only outside their own region.')]
def insert_after_heading(doc, prefix, blocks):
    ps = doc.paragraphs
    idx = [i for i, p in enumerate(ps) if p.style.name == 'Heading 1' and p.text.startswith(prefix)]
    if not idx: print('missing', prefix); return
    nxt = next((p for p in ps[idx[0] + 1:] if p.style.name == 'Heading 1'), None)
    for k, t in blocks:
        par = doc.add_paragraph(t, style={'h2': 'Heading 2', 'p': 'Normal', 'b': 'List Bullet'}[k])
        if nxt is not None: nxt._p.addprevious(par._p)
def apply(doc):
    insert_after_heading(doc, '4. For regional and assistant coordinators', RC); insert_after_heading(doc, '5. For HQ', HQ)
def bump(doc, a, b):
    for p in doc.paragraphs:
        if a in p.text:
            for r in p.runs: r.text = r.text.replace(a, b)
g = B + 'KF4_Field_Command_Centre_User_Guide.docx'; d = docx.Document(g); apply(d); bump(d, 'Version 2.1', 'Version 2.2'); d.save(g)
m = B + 'KF4_Field_Operations_Manual_INTERIM_v0.9.docx'; d = docx.Document(m); apply(d); bump(d, 'Version 0.9', 'Version 1.0')
for t in d.tables:
    if t.rows and 'ersion' in t.rows[0].cells[0].text:
        c = t.add_row().cells
        for cell, v in zip(c, ['1.0', '24 Oct 2026', 'Lead Coordinator (LC) role for GEP, KACODA and FAWOCHIWE: read access to every region of the organization; HQ sets it in Admin.']): cell.text = v
d.save(m); os.rename(m, B + 'KF4_Field_Operations_Manual_INTERIM_v1.0.docx')
print('ok')
