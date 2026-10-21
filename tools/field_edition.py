# -*- coding: utf-8 -*-
# Builds the FIELD EDITIONS of the guide, manual and quick start: no HQ, administration or confidential parts.
import docx, os, re
from docx.text.paragraph import Paragraph
B=os.path.expanduser('~/LearnImpact_ED/KiuFunza Analysis/KF4_Endline_2026_Build/'); O=B+'Field_Edition/'; os.makedirs(O,exist_ok=True)
def level(el,d):
    if not el.tag.endswith('}p'): return None
    p=Paragraph(el,d); n=p.style.name
    return 1 if n=='Heading 1' else 2 if n=='Heading 2' else None
def remove(d,start_pred,stop_levels,keep_heading=False,replace=None):
    els=list(d.element.body.iterchildren()); i=0; removed=0
    while i<len(els):
        el=els[i]
        if el.tag.endswith('}p') and start_pred(Paragraph(el,d)):
            lv=level(el,d); j=i+1
            while j<len(els) and not els[j].tag.endswith('}sectPr') and not (level(els[j],d) in stop_levels and level(els[j],d)<=lv): j+=1
            for k in range(i+(1 if keep_heading else 0),j): els[k].getparent().remove(els[k]); removed+=1
            if replace:
                par=d.add_paragraph(replace); (els[j] if j<len(els) else d.element.body[-1]).addprevious(par._p)
            i=j; continue
        i+=1
    return removed
H1=lambda pre:(lambda p:p.style.name=='Heading 1' and p.text.startswith(pre))
H2=lambda pre:(lambda p:p.style.name=='Heading 2' and p.text.startswith(pre))
PLACE='This section is for HQ only. HQ colleagues have the full edition of this guide in the Help page.'
# ---- user guide
g=docx.Document(B+'KF4_Field_Command_Centre_User_Guide.docx')
remove(g,H1('5. For HQ'),{1},keep_heading=True,replace=PLACE)
for p in g.paragraphs:
    for r in p.runs:
        if 'Version 1.4' in r.text: r.text=r.text.replace('Version 1.4','Version 1.4 · Field edition')
g.save(O+'KF4_Field_Command_Centre_User_Guide_FIELD.docx')
# ---- manual
m=docx.Document(B+'KF4_Field_Operations_Manual_INTERIM_v0.6.docx')
print('C5',remove(m,H1('5. For HQ'),{1},keep_heading=True,replace=PLACE))
print('F backup',remove(m,H2('Where the data is, and how it is backed up'),{1,2}))
print('G',remove(m,H1('Part G'),{1}))
print('hold',remove(m,H2('What we hold, and where'),{1,2}))
# the HQ-only line in the breach section, the payments line, and cover text
for p in list(m.paragraphs):
    if p.text.startswith('HQ acts in the system'): p._p.getparent().remove(p._p)
    elif p.text.startswith('Still to come: the payment logic'): p._p.getparent().remove(p._p)
    elif p.text.startswith('Version 0.2'):
        for r in p.runs: r.text=re.sub(r'Draft for review by Mundo and Joram, then for the field team','Field edition: HQ, administration and confidential parts are not included',r.text)
for p in m.paragraphs:
    if 'Part G' in p.text or 'Mundo and Joram review' in p.text:
        for r in p.runs: r.text=r.text.replace(' and the open decisions in Part G','').replace('Part G','the HQ edition')
for p in m.paragraphs:
    for r in p.runs:
        if 'Open decisions are listed' in r.text: r.text=r.text.replace('Open decisions are listed in the HQ edition with an owner. ','')
for t in m.tables:
    for row in list(t.rows):
        if 'payment' in ' '.join(c.text for c in row.cells).lower(): row._tr.getparent().remove(row._tr)
m.save(O+'KF4_Field_Operations_Manual_FIELD_v0.6.docx')
# ---- quick start: pages 1 to 3 only
q=docx.Document(B+'KF4_LI-CC_Quick_Start.docx'); els=list(q.element.body.iterchildren()); cut=None
for i,el in enumerate(els):
    if el.tag.endswith('}p') and Paragraph(el,q).text.startswith('For HQ'): cut=i; break
# remove from the page-break paragraph before the HQ title (two elements earlier) to the end
start=max(0,cut-2)
for el in els[start:]:
    if not el.tag.endswith('}sectPr'): el.getparent().remove(el)
q.save(O+'KF4_LI-CC_Quick_Start_FIELD.docx')
# check: nothing confidential left
bad=['HQ_ACCESS','passphrase','Passphrase','wrangler','GitHub','Cloudflare','AUTH_SALT','HQ_SECRET','decrypt','Open decisions','Mundo and Joram review','Admin > People','Guests (invited']
for f in os.listdir(O):
    d=docx.Document(O+f); txt='\n'.join(p.text for p in d.paragraphs)+'\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
    print(f,len(d.paragraphs),[b for b in bad if b in txt])
