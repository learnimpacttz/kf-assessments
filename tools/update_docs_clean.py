# -*- coding: utf-8 -*-
# Data cleaning page changes of 27 Oct 2026: guide 2.2 -> 2.3, manual 1.0 -> 1.1, HQ quick start gets one line.
import docx, os
B = os.path.expanduser('~/LearnImpact_ED/KiuFunza Analysis/KF4_Endline_2026_Build/')
HQ = [('h2', 'Correcting a pupil number, and what the Data cleaning page does for you (HQ)'),
 ('p', 'Data cleaning, Suggest a deletion or an edit, Edit records: choose the Students project, find the school, tick the records with the wrong number and choose Pupil number / ID. Beside each ticked record the page shows its current number and a box for the new one. Type only the pupil number, from 1 to 400. It is the last digits of the pupil ID, so for 2026PS21030253006 you type 6, not the whole ID. Beside the box the page shows what the ID will become, or what is wrong with what was typed. When two pupils share a number, tick both and keep one as it is: only the other one needs a new number. A request that would change nothing (the same number on every ticked record) is refused.'),
 ('b', 'What you type (new numbers, the value to change to, the reason, the confirm number and notes) stays in its box when the page refreshes and while you move around the form. It is cleared once the step has been sent, and kept when a step fails so you can correct it and press again.'),
 ('b', 'While a step waits for KoBo (up to a minute) the button shows Sending and the page asks you to wait. A second press is ignored, so one suggestion cannot be sent twice.'),
 ('b', 'A record that is already in someone\'s open request cannot be put in a second request. The row says whose request holds it and where it waits (for example Mundo\'s request, waiting for Joram) and has an Open it button that goes straight to that request. Joram confirms what Mundo suggests, Mundo gives the final decision, and the person who suggested it can withdraw it.'),
 ('b', 'What happens to a request (done, stopped, a wrong confirm number) is shown on that request\'s own card. Requests saved at the same moment by two people no longer erase each other; if two people change the same request at the same moment, the second is asked to open it again and repeat the step.')]
QS = 'Pupil number: type only the last digits of the ID (6 for 2026PS21030253006), and wait for Sending to finish. A record marked "in a request" is in someone\'s open request: press Open it and act on it there.'
def insert(doc):
    ps = doc.paragraphs
    idx = [i for i, p in enumerate(ps) if p.style.name == 'Heading 1' and p.text.startswith('5. For HQ')]
    if not idx: print('missing HQ section'); return
    nxt = next((p for p in ps[idx[0] + 1:] if p.style.name == 'Heading 1'), None)
    for k, t in HQ:
        par = doc.add_paragraph(t, style={'h2': 'Heading 2', 'p': 'Normal', 'b': 'List Bullet'}[k])
        if nxt is not None: nxt._p.addprevious(par._p)
def bump(doc, a, b):
    for p in doc.paragraphs:
        if a in p.text:
            for r in p.runs: r.text = r.text.replace(a, b).replace('24 October 2026', '27 October 2026')
g = B + 'KF4_Field_Command_Centre_User_Guide.docx'; d = docx.Document(g); insert(d); bump(d, 'Version 2.2', 'Version 2.3'); d.save(g)
m = B + 'KF4_Field_Operations_Manual_INTERIM_v1.0.docx'; d = docx.Document(m); insert(d); bump(d, 'Version 1.0', 'Version 1.1')
for t in d.tables:
    if t.rows and 'ersion' in t.rows[0].cells[0].text:
        c = t.add_row().cells
        for cell, v in zip(c, ['1.1', '27 Oct 2026', 'Data cleaning page: typed entries kept on refresh, clearer pupil-number rules, whose request holds a record, results on the request card, no double send.']): cell.text = v
d.save(m); os.rename(m, B + 'KF4_Field_Operations_Manual_INTERIM_v1.1.docx')
q = B + 'KF4_LI-CC_Quick_Start.docx'; d = docx.Document(q)
at = next((p for p in d.paragraphs if p.text.startswith('Never use the KoBo web editor')), None)
if at is None: print('quick start anchor missing')
else:
    par = d.add_paragraph(QS, style='List Bullet'); at._p.addnext(par._p); d.save(q)
print('ok')
