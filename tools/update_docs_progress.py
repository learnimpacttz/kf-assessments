# -*- coding: utf-8 -*-
# Adds the region progress board (RC and ARC) and the standing card (test admins) to the guide (2.0 -> 2.1) and manual (0.8 -> 0.9).
import docx, os
B = os.path.expanduser('~/LearnImpact_ED/KiuFunza Analysis/KF4_Endline_2026_Build/')
RC = [('h2', 'Region progress (Region page)'),
 ('p', 'The top of the Region page now shows where the region stands, against the plan, without leaving the page.'),
 ('b', 'Two counters: pupils tested and schools complete. Under them a bar shows every school by state: complete, in progress, behind plan, visit today, planned, no date yet. Tap a colour to list those schools. The plan line compares the schools due by the plan with the schools complete, by working day, with today marked.'),
 ('b', 'Four tiles: queries open (and how many are overdue), responses waiting for your review, questions waiting for you, and how clean the data is (the share of records not touched by an open check).'),
 ('b', 'School tracker: the eight schools that most need attention first (Show all lists every school). Each row shows the planned date and who leads, pupils by grade, progress and what is pending: visit, notices, sampling attendance, queries, questions. Tap a school for its full picture.'),
 ('b', 'In a school: what is pending, the plan with the notices to the ward officer, head teacher and team and the result of the visit, any plan changes, every query with the response, the reviews and the comments (including approved, closed and withdrawn ones), and the questions asked about the school and their answers.'),
 ('b', 'Three folds under the tracker, closed until you open them: Waiting on (everything open, grouped by who has the next step, red first), The team (who is carrying the team and who needs a hand, with the reasons) and RC and ARC: how active and how clean (responses written, reviews done and how fast, reviews and questions waiting, and the clean share in the schools each one leads).'),
 ('b', 'Use it to thank people and to help. A person can be in both lists. Nothing here changes a record, a payment or a query.')]
VOL = [('b', 'Your standing (My day and My stats): a short card with your place in the team on pupils tested, records sent the same day and test quality, and how many queries are waiting for you. My stats shows what you give the team, what is waiting on you, and colleagues who are carrying the team, so you can thank them.'),
 ('b', 'Your place is compared with the other test admins in your region who have tested at least 20 pupils. A query waits on you when no response has been written 2 working days after the visit. Colleagues who are behind are only counted, never named; your RC and ARC help them.')]
def insert_after_heading(doc, prefix, blocks):
    ps = doc.paragraphs
    idx = [i for i, p in enumerate(ps) if p.style.name == 'Heading 1' and p.text.startswith(prefix)]
    if not idx: print('missing', prefix); return
    nxt = next((p for p in ps[idx[0] + 1:] if p.style.name == 'Heading 1'), None)
    for k, t in blocks:
        par = doc.add_paragraph(t, style={'h2': 'Heading 2', 'p': 'Normal', 'b': 'List Bullet'}[k])
        if nxt is not None: nxt._p.addprevious(par._p)
def apply(doc):
    insert_after_heading(doc, '3. For test admins', VOL); insert_after_heading(doc, '4. For regional and assistant coordinators', RC)
def bump(doc, a, b, date_a='22 October 2026', date_b='24 October 2026'):
    for p in doc.paragraphs:
        if a in p.text:
            for r in p.runs: r.text = r.text.replace(a, b).replace(date_a, date_b)
g = B + 'KF4_Field_Command_Centre_User_Guide.docx'; d = docx.Document(g); apply(d); bump(d, 'Version 2.0', 'Version 2.1'); d.save(g)
m = B + 'KF4_Field_Operations_Manual_INTERIM_v0.8.docx'; d = docx.Document(m); apply(d); bump(d, 'Version 0.8', 'Version 0.9')
for t in d.tables:
    if t.rows and 'ersion' in t.rows[0].cells[0].text:
        c = t.add_row().cells
        for cell, v in zip(c, ['0.9', '24 Oct 2026', 'Region progress for RCs and ARCs (schools against the plan, school tracker and drawer, waiting on, the team, activity and cleanliness); standing card for test admins.']): cell.text = v
d.save(m); os.rename(m, B + 'KF4_Field_Operations_Manual_INTERIM_v0.9.docx')
print('ok')
