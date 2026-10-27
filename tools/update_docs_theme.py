# -*- coding: utf-8 -*-
# Dark mode readability and the by-day pupils table (27 Oct 2026): guide 2.4 -> 2.5, manual 1.2 -> 1.3.
import docx, os
B = os.path.expanduser('~/LearnImpact_ED/KiuFunza Analysis/KF4_Endline_2026_Build/')
EVERYONE = [('p', 'Appearance. In the account menu (your initials, top right), Appearance lets you choose light, dark or automatic (it follows your phone). Every page, chart, table and pop-up was checked for readable contrast in both light and dark. If anything is hard to read in either, send a screenshot through Questions.')]
HQ = [('h2', 'Pupils tested, by region and day'),
      ('p', 'On the Bridge, the table Pupils tested, by region and day has one row for each region and one column for each day with testing (the last 14 days). The stronger the color, the more pupils were tested that day. An empty or pale cell on a field day is a region to ask about. A very large day, such as the pilot day, does not wash out the other days. On a phone, slide the table sideways to see all the days. (It used to show whole weeks, which made a week with only two days look like a drop.)')]
def insert(doc, heading, items):
    ps = doc.paragraphs
    idx = [i for i, p in enumerate(ps) if p.style.name == 'Heading 1' and p.text.startswith(heading)]
    if not idx: print('missing', heading); return
    nxt = next((p for p in ps[idx[0] + 1:] if p.style.name == 'Heading 1'), None)
    for k, t in items:
        par = doc.add_paragraph(t, style={'h2': 'Heading 2', 'p': 'Normal', 'b': 'List Bullet'}[k])
        if nxt is not None: nxt._p.addprevious(par._p)
def bump(doc, a, b):
    for p in doc.paragraphs:
        if a in p.text:
            for r in p.runs: r.text = r.text.replace(a, b)
g = B + 'KF4_Field_Command_Centre_User_Guide.docx'; d = docx.Document(g)
insert(d, '2. Getting started (everyone)', EVERYONE); insert(d, '5. For HQ', HQ); bump(d, 'Version 2.4', 'Version 2.5'); d.save(g)
m = B + 'KF4_Field_Operations_Manual_INTERIM_v1.2.docx'; d = docx.Document(m)
insert(d, '2. Getting started (everyone)', EVERYONE); insert(d, '5. For HQ', HQ); bump(d, 'Version 1.2', 'Version 1.3')
for t in d.tables:
    if t.rows and 'ersion' in t.rows[0].cells[0].text:
        c = t.add_row().cells
        for cell, v in zip(c, ['1.3', '27 Oct 2026', 'Dark mode readability fixed on every page; the Bridge pupils table now shows days, not weeks.']): cell.text = v
d.save(m); os.rename(m, B + 'KF4_Field_Operations_Manual_INTERIM_v1.3.docx')
print('ok')
