# -*- coding: utf-8 -*-
# Adds answer patterns and phones (Team patterns) to the HQ section of the guide (1.9 -> 2.0) and manual (0.7 -> 0.8).
import docx, os
B = os.path.expanduser('~/LearnImpact_ED/KiuFunza Analysis/KF4_Endline_2026_Build/')
HQ = [('h2', 'Answer patterns and phones (Team patterns)'),
 ('p', 'Team patterns now also looks at the answers themselves and at the phones a person uses. Both are checks of the facts, so one is enough to put a person on the list. Neither changes a record or a payment.'),
 ('b', 'Answers that look alike. A test is made of blocks of five questions, and a block stops once three are correct. A pupil between two levels gives mixed blocks (some right, some wrong). A tester who asks only the first question and marks the rest the same way almost never gives a mixed block. The page compares each person with the other testers at the same schools and grades, so a strong or weak class does not look like a pattern. A person is flagged with at least 30 full tests when their mixed blocks are at 65% or less of what the others got and at least 40% of their full tests have no mixed block. When the gap is very large (mixed blocks at 40% or less, no mixed block in 60% or more of tests, the same result in every block in 40% or more) the person goes to Needs attention.'),
 ('b', 'Open Details to see the exact submissions (day, school, grade, minutes, start time and KoBo id) to open in KoBo, or to re-test a few of those pupils. Pupil names and numbers are never shown. Only counts are kept in the system, never the answers.'),
 ('b', 'Phones. The page flags two phones used by one name at different schools on the same day, a phone shared with someone else for a quarter or more of a person\'s tests, a main phone that keeps changing, and two or more phones on two days or more. The pilot day is not counted. Two phones at different schools on one day shows even before the minimum days and tests, because one person cannot be at both schools.'),
 ('b', 'What to do. Start with a supportive conversation. Look at a few recordings or sit in on a visit, and ask a colleague to re-test a small sample of pupils from the listed tests (a back-check). Possible reasons are shown on each card, for example a replaced or lent phone. If the data itself is in doubt, raise it through the Queries chain; nothing is deleted or changed from this page. Patterns are for HQ only and Advisors do not see them.'),
 ('b', 'The counters for answers are built when a record is read. After the first release, records already read before it carry no answer counts until the Students data is rebuilt, which the Executive Director or the systems owner approves.')]
def insert_after_heading(doc, prefix, blocks):
    ps = doc.paragraphs
    idx = [i for i, p in enumerate(ps) if p.style.name == 'Heading 1' and p.text.startswith(prefix)]
    if not idx: print('missing', prefix); return
    nxt = next((p for p in ps[idx[0] + 1:] if p.style.name == 'Heading 1'), None)
    for k, t in blocks:
        par = doc.add_paragraph(t, style={'h2': 'Heading 2', 'p': 'Normal', 'b': 'List Bullet'}[k])
        if nxt is not None: nxt._p.addprevious(par._p)
def bump(doc, a, b):
    for p in doc.paragraphs:
        if a in p.text:
            for r in p.runs: r.text = r.text.replace(a, b)
g = B + 'KF4_Field_Command_Centre_User_Guide.docx'; d = docx.Document(g); insert_after_heading(d, '5. For HQ', HQ); bump(d, 'Version 1.9', 'Version 2.0'); d.save(g)
m = B + 'KF4_Field_Operations_Manual_INTERIM_v0.7.docx'; d = docx.Document(m); insert_after_heading(d, '5. For HQ', HQ); bump(d, 'Version 0.7', 'Version 0.8')
for t in d.tables:
    if t.rows and 'ersion' in t.rows[0].cells[0].text:
        c = t.add_row().cells
        for cell, v in zip(c, ['0.8', '22 Oct 2026', 'Team patterns: answers that look alike within a block, and phones (two at different schools on one day, shared, changing).']): cell.text = v
d.save(m); os.rename(m, B + 'KF4_Field_Operations_Manual_INTERIM_v0.8.docx')
print('ok')
